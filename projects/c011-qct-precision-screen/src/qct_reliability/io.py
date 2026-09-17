from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook


DATA_SHEETS = ("AL to MO", "MT to WY & PR")


def read_hud_workbook(path: Path, sheets: Iterable[str] = DATA_SHEETS) -> pd.DataFrame:
    """Read values without modifying or recalculating the official workbook."""
    book = load_workbook(path, read_only=True, data_only=True)
    frames: list[pd.DataFrame] = []
    for name in sheets:
        rows = book[name].iter_rows(values_only=True)
        header = list(next(rows))
        frames.append(pd.DataFrame(rows, columns=header))
    return pd.concat(frames, ignore_index=True)
