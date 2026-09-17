from pathlib import Path
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
out = {}
for path in sorted((ROOT / "data" / "raw").glob("qct_data_*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=False)
    out[path.name] = {
        "sheets": [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                    "state": ws.sheet_state} for ws in book.worksheets],
        "defined_names": sorted(book.defined_names),
    }
target = ROOT / "outputs" / "diagnostics" / "workbook_structure.json"
target.write_text(json.dumps(out, indent=2) + "\n", encoding="utf-8")
print(target)
